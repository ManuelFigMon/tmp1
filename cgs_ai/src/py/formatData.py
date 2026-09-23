"""
=====================================================================
  Program Name  : formatData.py
  Author        : Manuel Figallo
  Purpose       : Convert a CSV *or an Excel workbook* into a styled Excel
                  workbook with a SAS ODS-style look and feel, suitable
                  for distribution.
  Version       : 1.0beta
  Created       : 2026-08-26
  Last Modified : 2026-09-06

  Dependencies:
    openpyxl, imported lazily, for WRITING. Reading an .xlsx needs nothing
    beyond the standard library (src/utils/xlsx.py).

  Description:
    The default "corporate" style is a navy title banner, blue header row,
    zebra striping and a frozen, auto-filtered header -- clean and
    business-appropriate, matching SAS ODS output.

    Renamed from formatCSV, which now accepts an .xlsx as well as a .csv --
    "CSV" had stopped describing what it does. formatCSV remains as an alias
    so existing pipelines and training material keep working.

    FEEDING BACK ITS OWN OUTPUT. A workbook this function produces has the
    title banner on row 1 and the headers on row 2. Reading one back detects
    that automatically, so the output of one call can be the input of the
    next with no extra parameter. Pass HeaderRow to override the detection.

  Input Parameters (required first):
    InputPath       (REQUIRED, str) - source .csv or .xlsx.
    OutputExcelPath (REQUIRED, str) - destination .xlsx.
    FormatType      (optional, str, default "corporate") - corporate |
                      corporatev2 | plain | minimal.
    SheetName       (optional, str, default "Report") - the OUTPUT sheet.
    Title           (optional, str) - banner text; defaults to the file name.
    InputSheet      (optional, str|int) - which worksheet to read when the
                      input is .xlsx; default the first.
    HeaderRow       (optional, int, default 0 = detect) - 1-based header row
                      in an .xlsx input.
    InputCsvPath    (deprecated) - the old name for InputPath.
=====================================================================
"""

from __future__ import annotations

import io
import sys
from pathlib import Path
from typing import Any, Dict

_HERE = Path(__file__).resolve()
if str(_HERE.parent.parent.parent) not in sys.path:
    sys.path.insert(0, str(_HERE.parent.parent.parent))

from src.utils.helpers import ensureParent, readCsv  # noqa: E402
from src.utils.logger import logInfo                 # noqa: E402
from src.utils.xlsx import readSheet                 # noqa: E402

__version__ = "1.0beta"

#: Palette per format type: (banner, header, stripe, header font colour).
FORMAT_STYLES: Dict[str, Dict[str, str]] = {
    "corporate":   {"banner": "1F3864", "header": "2E75B6", "stripe": "DCE6F1",
                    "headerFont": "FFFFFF", "bannerFont": "FFFFFF"},
    "corporatev2": {"banner": "1F3864", "header": "2E75B6", "stripe": "DCE6F1",
                    "headerFont": "FFFFFF", "bannerFont": "FFFFFF"},
    "plain":       {"banner": "FFFFFF", "header": "D9D9D9", "stripe": "FFFFFF",
                    "headerFont": "000000", "bannerFont": "000000"},
    "minimal":     {"banner": "FFFFFF", "header": "FFFFFF", "stripe": "FFFFFF",
                    "headerFont": "000000", "bannerFont": "000000"},
}
#: Format types that get zebra striping on the data rows.
STRIPED_FORMATS = frozenset({"corporate", "corporatev2"})
DEFAULT_FORMAT_TYPE = "corporate"
MAX_COLUMN_WIDTH = 60


#: Input extensions read as a workbook rather than as delimited text.
EXCEL_EXTENSIONS = frozenset({".xlsx", ".xlsm"})


def readTable(InputPath: str, InputSheet: Any = None,
              HeaderRow: int = 0) -> list:
    """Read the input, whichever of the two shapes it is.

    Parameters:
        InputPath (str)      - a .csv (or any delimited text) or a .xlsx.
        InputSheet (str|int) - worksheet to read from an .xlsx; default first.
        HeaderRow (int)      - 1-based header row in an .xlsx; 0 detects it.
    Returns:
        list[dict] keyed by column name -- the same shape either way, so the
        formatter below does not care where the rows came from.
    Raises:
        OSError / KeyError - the file cannot be read, or the sheet is absent.
    """
    if Path(InputPath).suffix.lower() in EXCEL_EXTENSIONS:
        return readSheet(InputPath, sheet=InputSheet, headerRow=HeaderRow)
    return readCsv(InputPath)


def formatData(InputPath: str = "", OutputExcelPath: str = "",
               FormatType: str = DEFAULT_FORMAT_TYPE,
               SheetName: str = "Report",
               Title: str = "",
               InputSheet: Any = None,
               HeaderRow: int = 0,
               InputCsvPath: str = "") -> Dict[str, Any]:
    """Render a CSV or an Excel workbook as a styled Excel workbook.

    Parameters:
        InputPath (str)       - REQUIRED source .csv or .xlsx.
        OutputExcelPath (str) - REQUIRED destination .xlsx.
        FormatType (str)      - corporate (default) | corporatev2 |
                                plain | minimal. "corporate" and
                                "corporatev2" are a navy banner, blue header
                                and zebra striping; they render identically,
                                with corporatev2 reserved as the versioned
                                name for that look.
        SheetName (str)       - OUTPUT worksheet name; default "Report".
        Title (str)           - banner text; defaults to the input filename.
        InputSheet (str|int)  - which worksheet to read from an .xlsx input;
                                default the first.
        HeaderRow (int)       - 1-based header row in an .xlsx input; 0
                                (default) detects it, which is what lets a
                                workbook this function wrote be read back
                                without counting past its title banner.
        InputCsvPath (str)    - DEPRECATED alias for InputPath, kept so
                                existing pipelines keep running.
    Returns:
        dict with OutputExcelPath, RowCount, ColumnCount, FormatType and
        InputPath.
    Raises:
        ValueError  - a required parameter is missing or FormatType unknown.
        ImportError - openpyxl is not installed (message says how to fix).
        OSError     - the input cannot be read.

    Use in claims processing:
        Turn a raw scan or claims extract into a report an analyst or manager
        can open directly, without hand-formatting in Excel each cycle -- and
        re-style a workbook a colleague sent without converting it to CSV
        first.
    """
    InputPath = str(InputPath or InputCsvPath or "").strip()
    if not InputPath:
        raise ValueError("required parameter 'InputPath' is missing or empty")
    if not OutputExcelPath or not str(OutputExcelPath).strip():
        raise ValueError("required parameter 'OutputExcelPath' is missing or empty")
    if FormatType not in FORMAT_STYLES:
        raise ValueError(f"unknown FormatType {FormatType!r}; expected one of: "
                         f"{', '.join(FORMAT_STYLES)}")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError(
            "formatData requires openpyxl for Excel styling. Install it with "
            "'pip install openpyxl'.")

    rows = readTable(InputPath, InputSheet=InputSheet, HeaderRow=HeaderRow)
    columns = list(rows[0].keys()) if rows else []
    style = FORMAT_STYLES[FormatType]
    banner = Title or Path(InputPath).stem

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SheetName

    # Row 1: title banner spanning the table width.
    sheet.cell(row=1, column=1, value=banner)
    cell = sheet.cell(row=1, column=1)
    cell.font = Font(bold=True, size=14, color=style["bannerFont"])
    cell.fill = PatternFill("solid", fgColor=style["banner"])
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26
    if columns:
        sheet.merge_cells(start_row=1, start_column=1,
                          end_row=1, end_column=len(columns))

    # Row 2: column headers.
    for index, column in enumerate(columns, start=1):
        header = sheet.cell(row=2, column=index, value=column)
        header.font = Font(bold=True, color=style["headerFont"])
        header.fill = PatternFill("solid", fgColor=style["header"])
        header.alignment = Alignment(vertical="center", wrap_text=True)

    # Rows 3+: data with zebra striping.
    stripe = PatternFill("solid", fgColor=style["stripe"])
    for rowIndex, row in enumerate(rows, start=3):
        for colIndex, column in enumerate(columns, start=1):
            dataCell = sheet.cell(row=rowIndex, column=colIndex,
                                  value=row.get(column, ""))
            if FormatType in STRIPED_FORMATS and rowIndex % 2 == 1:
                dataCell.fill = stripe

    if columns:
        sheet.freeze_panes = "A3"
        sheet.auto_filter.ref = (f"A2:{get_column_letter(len(columns))}"
                                 f"{max(2, len(rows) + 2)}")
        for index, column in enumerate(columns, start=1):
            widest = max([len(str(column))] +
                         [len(str(r.get(column, ""))) for r in rows[:200]] or [10])
            sheet.column_dimensions[get_column_letter(index)].width = \
                min(MAX_COLUMN_WIDTH, max(10, widest + 2))

    ensureParent(OutputExcelPath)
    # Save via BytesIO: the Snowflake workspace filesystem does not support
    # seek() on writable files, which openpyxl's ZIP writer requires. Writing
    # to an in-memory buffer and flushing it bypasses that limit, and behaves
    # identically everywhere else.
    buffer = io.BytesIO()
    workbook.save(buffer)
    with open(OutputExcelPath, "wb") as handle:
        handle.write(buffer.getvalue())
    logInfo(f"formatted {len(rows)} row(s) x {len(columns)} column(s) "
            f"[{FormatType}] {Path(InputPath).name} -> {OutputExcelPath}")
    return {"OutputExcelPath": OutputExcelPath, "RowCount": len(rows),
            "ColumnCount": len(columns), "FormatType": FormatType,
            "InputPath": InputPath}


def formatCSV(InputCsvPath: str = "", OutputExcelPath: str = "",
              FormatType: str = DEFAULT_FORMAT_TYPE,
              SheetName: str = "Report", Title: str = "",
              **kwargs: Any) -> Dict[str, Any]:
    """DEPRECATED alias for formatData, kept so existing callers keep working.

    Parameters: as formatData, with the original parameter name InputCsvPath.
    Returns: whatever formatData returns.

    Not removed, and not warned about on every call: this name is in running
    pipelines, in the SAS wrappers and in the training handouts. It reads
    .xlsx too, because it is the same function underneath.
    """
    return formatData(InputPath=InputCsvPath, OutputExcelPath=OutputExcelPath,
                      FormatType=FormatType, SheetName=SheetName, Title=Title,
                      **kwargs)
